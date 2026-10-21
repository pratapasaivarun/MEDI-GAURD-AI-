from decimal import Decimal
from io import BytesIO

from docx import Document
import fitz

import mediguard.ai.agents as agents
from mediguard.engine.models import CategoryComparison, Clause, ComparisonResult, SettlementData
from mediguard.reporting.reports import (
    ESTIMATE_DISCLAIMER,
    build_appeal_letter_pdf,
    build_comparison_appeal_letter_docx,
    build_comparison_appeal_letter_pdf,
    build_comparison_appeal_letter_text,
    build_hospital_note_docx,
    build_hospital_note_pdf,
    build_hospital_note_text,
)


def _comparison():
    clause = Clause("Clause 3.4: Expenses incurred during hospitalization, including medicines, are payable subject to co-payment.", 9, "3.4")
    row = CategoryComparison("PHARMACY", Decimal("10000"), Decimal("0"), Decimal("10000"), "UNEXPLAINED_DEDUCTION", "Medicines outside the hospitalization window", "No rule supports a deduction in this category.", clause, Decimal("9000"))
    return ComparisonResult(
        Decimal("20400"), Decimal("29400"), Decimal("9000"), [row], "possible_discrepancy", [],
        Decimal("132000"), Decimal("102600"), Decimal("111600"),
    )


def test_appeal_letter_uses_expected_numbers_and_neutral_language():
    comparison = _comparison()
    settlement = SettlementData("CLM-2026-0417", Decimal("102600"), [], Decimal("11400"), False, [], Decimal("132000"))
    letter = build_comparison_appeal_letter_text({
        "claim_number": "CLM-2026-0417", "patient_name": "Ravi Reddy", "hospital_name": "Sunrise Multispeciality Hospital", "policy_number": "POL-558201", "stay_dates": "10 to 14 Aug 2026",
    }, comparison, settlement)
    for value in ("₹1,32,000", "₹1,02,600", "₹1,11,600", "₹9,000", "₹10,000", "Clause 3.4", "page 9"):
        assert value in letter
    assert ESTIMATE_DISCLAIMER in letter
    assert "Insurer reason: Medicines outside the hospitalization window. Review point:" in letter
    for forbidden in ("fraud", "illegal", "cheat", "wrongful", "intentional"):
        assert forbidden not in letter.lower()


def test_hospital_note_states_undisputed_amount_and_disclaimer():
    note = build_hospital_note_text({"claim_number": "CLM-2026-0417", "patient_name": "Ravi Reddy"}, _comparison())
    assert "under insurer review" in note
    assert "₹9,000" in note
    assert "₹20,400" in note
    assert ESTIMATE_DISCLAIMER in note


def test_pdf_and_docx_report_builders_emit_valid_container_signatures():
    comparison = _comparison()
    settlement = SettlementData("CLM-2026-0417", Decimal("102600"), [], Decimal("11400"), False, [], Decimal("132000"))
    claim = {"claim_number": "CLM-2026-0417", "patient_name": "Ravi Reddy"}
    assert build_comparison_appeal_letter_pdf(claim, comparison, settlement).startswith(b"%PDF")
    assert build_comparison_appeal_letter_docx(claim, comparison, settlement).startswith(b"PK")
    assert build_hospital_note_pdf(claim, comparison).startswith(b"%PDF")
    assert build_hospital_note_docx(claim, comparison).startswith(b"PK")


def test_provisional_appeal_is_marked_in_text_pdf_and_word_but_confirmed_is_not():
    comparison = _comparison()
    settlement = SettlementData("CLM-2026-0417", Decimal("102600"), [], Decimal("11400"), False, [], Decimal("132000"))
    claim = {"claim_number": "CLM-2026-0417", "patient_name": "Ravi Reddy"}

    provisional_text = build_comparison_appeal_letter_text(claim, comparison, settlement, provisional=True)
    confirmed_text = build_comparison_appeal_letter_text(claim, comparison, settlement)
    assert provisional_text.startswith("PROVISIONAL DRAFT")
    assert "PROVISIONAL DRAFT" not in confirmed_text

    provisional_pdf = fitz.open(stream=build_comparison_appeal_letter_pdf(claim, comparison, settlement, provisional=True), filetype="pdf")
    confirmed_pdf = fitz.open(stream=build_comparison_appeal_letter_pdf(claim, comparison, settlement), filetype="pdf")
    assert "PROVISIONAL DRAFT" in provisional_pdf[0].get_text()
    assert "PROVISIONAL DRAFT" not in confirmed_pdf[0].get_text()

    provisional_docx = Document(BytesIO(build_comparison_appeal_letter_docx(claim, comparison, settlement, provisional=True)))
    confirmed_docx = Document(BytesIO(build_comparison_appeal_letter_docx(claim, comparison, settlement)))
    assert "PROVISIONAL DRAFT" in "\n".join(paragraph.text for paragraph in provisional_docx.paragraphs)
    assert "PROVISIONAL DRAFT" not in "\n".join(paragraph.text for paragraph in confirmed_docx.paragraphs)

    provisional_hospital_text = build_hospital_note_text(claim, comparison, provisional=True)
    assert provisional_hospital_text.startswith("PROVISIONAL DRAFT")
    hospital_pdf = fitz.open(stream=build_hospital_note_pdf(claim, comparison, provisional=True), filetype="pdf")
    assert "PROVISIONAL DRAFT" in hospital_pdf[0].get_text()
    hospital_docx = Document(BytesIO(build_hospital_note_docx(claim, comparison, provisional=True)))
    assert "PROVISIONAL DRAFT" in "\n".join(paragraph.text for paragraph in hospital_docx.paragraphs)

    decision_pdf = fitz.open(stream=build_appeal_letter_pdf(claim, {"status": "approved", "payable_amount": 8100, "policy_terms_review_status": "provisional"}), filetype="pdf")
    assert "PROVISIONAL DRAFT" in decision_pdf[0].get_text()


def test_fabricated_llm_number_falls_back_to_fixed_explanation(monkeypatch):
    comparison = _comparison()
    fallback = agents.comparison_template_explanation(comparison)
    monkeypatch.setattr(agents, "_ollama_json", lambda *args, **kwargs: {"explanation": "The possible difference is INR 99999."})
    explanation = agents.explain_comparison(comparison)
    assert explanation == fallback
    assert "99999" not in explanation
